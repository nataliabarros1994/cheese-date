"""Automacao de geracao de planilhas Excel formatadas."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows

from cheesedate.logger import get_logger
from cheesedate.paths import DATA_OUTPUT

log = get_logger(__name__)

# Estilos reutilizaveis
FONTE_TITULO = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
FONTE_CABECALHO = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FUNDO_TITULO = PatternFill("solid", fgColor="E8A33D")
FUNDO_CABECALHO = PatternFill("solid", fgColor="6B8E4E")
CENTRO = Alignment(horizontal="center", vertical="center")


def _escrever_titulo(ws, texto: str, n_colunas: int) -> None:
    """Escreve um titulo mesclado no topo da aba."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_colunas)
    celula = ws.cell(row=1, column=1, value=texto)
    celula.font = FONTE_TITULO
    celula.fill = FUNDO_TITULO
    celula.alignment = CENTRO
    ws.row_dimensions[1].height = 28


def _escrever_dataframe(ws, df: pd.DataFrame, linha_inicial: int = 3) -> int:
    """Escreve um DataFrame na aba com cabecalho formatado.

    Retorna o numero da ultima linha escrita.
    """
    linha = linha_inicial
    for i, nome_coluna in enumerate(df.columns, start=1):
        celula = ws.cell(row=linha, column=i, value=str(nome_coluna))
        celula.font = FONTE_CABECALHO
        celula.fill = FUNDO_CABECALHO
        celula.alignment = CENTRO

    # Dados
    for linha_dados in dataframe_to_rows(df, index=False, header=False):
        linha += 1
        for i, valor in enumerate(linha_dados, start=1):
            ws.cell(row=linha, column=i, value=valor)

    # Ajusta largura das colunas
    for i, nome_coluna in enumerate(df.columns, start=1):
        largura = max(len(str(nome_coluna)), 14) + 2
        ws.column_dimensions[get_column_letter(i)].width = largura

    return linha


def _aba_kpis(wb: Workbook, kpis: dict) -> None:
    """Cria a aba de resumo com os KPIs principais."""
    ws = wb.active
    ws.title = "Resumo"
    _escrever_titulo(ws, "CheeseData - Resumo da Producao", 2)

    rotulos = {
        "total_registros": "Total de registros",
        "kg_total": "Kg total produzido",
        "litros_total": "Litros de leite usados",
        "faturamento_total": "Faturamento total (R$)",
        "tipos_diferentes": "Tipos de queijo",
        "rendimento_medio": "Rendimento medio",
    }
    linha = 3
    for chave, valor in kpis.items():
        ws.cell(row=linha, column=1, value=rotulos.get(chave, chave)).font = Font(bold=True)
        ws.cell(row=linha, column=2, value=valor)
        linha += 1

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20
    log.info("Aba 'Resumo' criada")


def _aba_com_grafico(wb: Workbook, nome: str, titulo: str, df: pd.DataFrame,
                     col_categoria: int, col_valor: int) -> None:
    """Cria uma aba com tabela de dados e um grafico de barras nativo do Excel."""
    ws = wb.create_sheet(nome)
    _escrever_titulo(ws, titulo, len(df.columns))
    ultima_linha = _escrever_dataframe(ws, df)

    # Grafico de barras nativo do Excel
    chart = BarChart()
    chart.title = titulo
    chart.type = "col"
    dados = Reference(ws, min_col=col_valor, min_row=3, max_row=ultima_linha)
    categorias = Reference(ws, min_col=col_categoria, min_row=4, max_row=ultima_linha)
    chart.add_data(dados, titles_from_data=True)
    chart.set_categories(categorias)
    chart.height = 9
    chart.width = 16
    ws.add_chart(chart, f"A{ultima_linha + 3}")
    log.info("Aba '%s' criada com grafico", nome)


def gerar_planilha_excel(kpis: dict, df_mensal: pd.DataFrame,
                         df_ranking: pd.DataFrame, df_consumo: pd.DataFrame,
                         df_dados: pd.DataFrame) -> str:
    """Gera a planilha Excel completa com todas as abas.

    Returns:
        Caminho do arquivo .xlsx gerado.
    """
    log.info("Gerando planilha Excel")
    wb = Workbook()

    # Aba 1 - KPIs
    _aba_kpis(wb, kpis)

    # Aba 2 - Producao mensal (grafico: coluna 1=mes, coluna 2=kg)
    _aba_com_grafico(wb, "Producao Mensal", "Producao Mensal de Queijos",
                     df_mensal, col_categoria=1, col_valor=2)

    # Aba 3 - Ranking de queijos (coluna 1=tipo, coluna 3=faturamento)
    _aba_com_grafico(wb, "Ranking Queijos", "Faturamento por Tipo de Queijo",
                     df_ranking, col_categoria=1, col_valor=3)

    # Aba 4 - Consumo de materia-prima
    _aba_com_grafico(wb, "Consumo Leite", "Consumo de Leite por Tipo",
                     df_consumo, col_categoria=1, col_valor=2)

    # Aba 5 - Dados completos (sem grafico)
    ws_dados = wb.create_sheet("Dados Completos")
    _escrever_titulo(ws_dados, "Registros de Producao", len(df_dados.columns))
    _escrever_dataframe(ws_dados, df_dados)
    log.info("Aba 'Dados Completos' criada")

    caminho = DATA_OUTPUT / "relatorio_producao_queijos.xlsx"
    wb.save(caminho)
    log.info("Planilha salva em %s", caminho)
    return str(caminho)
